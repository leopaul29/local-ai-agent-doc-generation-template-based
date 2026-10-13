import pytest
from docx import Document
from fastapi import HTTPException
from fastapi.testclient import TestClient
from openpyxl import Workbook

import app as server
from documents import fill_template, read_template


def make_template(path):
    """Write a template with {{ name }} and {{ notes }} fields in the format of path."""
    if path.suffix == ".docx":
        doc = Document()
        doc.add_paragraph("Hello {{ name }}")
        doc.add_table(rows=1, cols=1).cell(0, 0).text = "{{ notes }}"
        doc.sections[0].footer.paragraphs[0].text = "Footer {{ name }}"
        doc.save(path)
    elif path.suffix == ".xlsx":
        book = Workbook()
        book.active.append(["Name", "{{ name }}"])
        book.active.append(["Notes", "{{ notes }}", "=1+1"])
        book.save(path)
    else:
        path.write_text("Hello {{ name }}\n{{ notes }}\n", encoding="utf-8")


@pytest.mark.parametrize("ext", [".docx", ".xlsx", ".md", ".txt"])
def test_read_and_fill(tmp_path, ext):
    src, dst = tmp_path / f"template{ext}", tmp_path / f"filled{ext}"
    make_template(src)
    assert read_template(src)[1] == ["name", "notes"]

    fill_template(src, dst, {"name": "Ada & <Co>", "notes": "line 1\nline 2"})

    text, fields = read_template(dst)
    assert fields == []
    assert "Ada & <Co>" in text and "line 1" in text and "line 2" in text


def test_chat_flow(tmp_path, monkeypatch):
    for folder in ("TEMPLATES", "OUTPUTS"):
        (tmp_path / folder).mkdir()
        monkeypatch.setattr(server, folder, tmp_path / folder)
    (tmp_path / "secret.md").write_text("{{ name }}")
    monkeypatch.setattr(server, "ask_llm", lambda system, messages, fields: {"values": {"name": "Ada"}, "reply": "Notes?"})
    client = TestClient(server.app)

    assert client.post("/api/templates", files={"file": ("a.pdf", b"x")}).status_code == 400
    assert client.post("/api/templates", files={"file": ("a.md", b"no fields")}).status_code == 400
    template = client.post("/api/templates", files={"file": ("a.md", b"Hi {{ name }}. {{ notes }}")}).json()
    assert template["fields"] == ["name", "notes"]

    messages = [{"role": "user", "content": "I am Ada"}]
    answer = client.post("/api/chat", json={"template_id": template["id"], "messages": messages}).json()
    assert answer == {"reply": "Notes?", "values": {"name": "Ada", "notes": ""}}
    assert client.get(f"/api/download/{template['id']}").text == "Hi Ada. "

    traversal = client.post("/api/chat", json={"template_id": "../secret.md", "messages": messages})
    assert traversal.status_code == 404


def test_ollama_unreachable(monkeypatch):
    monkeypatch.setattr(server, "OLLAMA_URL", "http://127.0.0.1:9")
    with pytest.raises(HTTPException) as error:
        server.ask_llm("system", [], ["name"])
    assert error.value.status_code == 502
