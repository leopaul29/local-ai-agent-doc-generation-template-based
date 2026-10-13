"""Read and fill document templates. Placeholders use Jinja syntax: {{ field_name }}."""
from pathlib import Path

from docx.oxml.ns import qn
from docxtpl import DocxTemplate
from jinja2 import meta
from jinja2.sandbox import SandboxedEnvironment
from openpyxl import load_workbook

FORMATS = {".docx", ".xlsx", ".md", ".txt"}
# Sandboxed because templates are uploaded files: they must not run arbitrary Python.
TEXT_ENV = SandboxedEnvironment(keep_trailing_newline=True)
DOCX_ENV = SandboxedEnvironment(autoescape=True)  # .docx is XML: escape &, < and > in values


def read_template(path: Path) -> tuple[str, list[str]]:
    """Return the template text (context for the model) and its field names."""
    if path.suffix == ".docx":
        doc = DocxTemplate(path)
        paragraphs = doc.get_docx().element.body.iter(qn("w:p"))
        text = "\n".join("".join(t.text or "" for t in p.iter(qn("w:t"))) for p in paragraphs)
        return text, sorted(doc.get_undeclared_template_variables(DOCX_ENV))
    if path.suffix == ".xlsx":
        rows = (row for sheet in load_workbook(path) for row in sheet.iter_rows(values_only=True))
        text = "\n".join("\t".join(str(v) for v in row if v is not None) for row in rows)
    else:
        text = path.read_text(encoding="utf-8")
    return text, sorted(meta.find_undeclared_variables(TEXT_ENV.parse(text)))


def fill_template(src: Path, dst: Path, values: dict[str, str]) -> None:
    """Write src to dst with its fields replaced by values, keeping the format and styling."""
    if src.suffix == ".docx":
        doc = DocxTemplate(src)
        doc.render(values, DOCX_ENV, autoescape=True)
        doc.save(dst)
    elif src.suffix == ".xlsx":
        book = load_workbook(src)
        for sheet in book:
            for row in sheet.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and "{" in cell.value:
                        # shortcut: values are written as text, convert them if formulas must compute on them
                        cell.value = TEXT_ENV.from_string(cell.value).render(values)
        book.save(dst)
    else:
        dst.write_text(TEXT_ENV.from_string(src.read_text(encoding="utf-8")).render(values), encoding="utf-8")
